"""Fixtures that build real files rather than mocking the readers.

A test that mocks pypdf proves that the mock works. These build an actual PDF
and an actual .docx so the extraction path, including its page and section
addressing, is the one that runs.
"""

from __future__ import annotations

import subprocess
from pathlib import Path

import pytest


def build_pdf(pages: list[str]) -> bytes:
    """A minimal, uncompressed PDF with one text-bearing page per entry."""
    objects: list[str] = []
    kids = " ".join(f"{4 + 2 * index} 0 R" for index in range(len(pages)))
    objects.append("<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    for index, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {5 + 2 * index} 0 R >>"
        )
        objects.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")

    out = b"%PDF-1.4\n"
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return out


@pytest.fixture
def pdf_file(tmp_path: Path):
    def _make(pages: list[str], name: str = "runbook.pdf") -> Path:
        path = tmp_path / name
        path.write_bytes(build_pdf(pages))
        return path

    return _make


@pytest.fixture
def docx_file(tmp_path: Path):
    def _make(sections: list[tuple[str, str]], name: str = "manual.docx") -> Path:
        import docx

        document = docx.Document()
        for heading, body in sections:
            document.add_heading(heading, level=1)
            document.add_paragraph(body)
        path = tmp_path / name
        document.save(str(path))
        return path

    return _make


@pytest.fixture
def git_repo(tmp_path: Path):
    """A real local repository, so the repo guards run against real entries."""

    def _make(files: dict[str, str], name: str = "repo") -> tuple[Path, str]:
        root = tmp_path / name
        root.mkdir()
        for relative, content in files.items():
            target = root / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(content, encoding="utf-8")

        env = {"GIT_CONFIG_NOSYSTEM": "1", "PATH": _path()}
        run = lambda *args: subprocess.run(  # noqa: E731
            ["git", *args], cwd=str(root), check=True, capture_output=True, env=env
        )
        run("init", "--quiet", "-b", "main")
        run("config", "user.email", "ci@example.invalid")
        run("config", "user.name", "ci")
        run("add", "-A")
        run("commit", "--quiet", "-m", "fixture")
        head = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=str(root),
            check=True,
            capture_output=True,
            text=True,
            env=env,
        ).stdout.strip()
        return root, head

    return _make


def _path() -> str:
    import os

    return os.environ.get("PATH", "")
